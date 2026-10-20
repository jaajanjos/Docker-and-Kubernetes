# Gera a base fictícia de e-commerce (Base_Ficticia_ECommerce.xlsx) na mesma pasta deste script.
# Requisito: pip install openpyxl  |  Uso: python gerar_base.py
import random, datetime as dt
from collections import defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.utils import get_column_letter

random.seed(42)
import os
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Base_Ficticia_ECommerce.xlsx")
F = "Arial"
HDR = PatternFill("solid", fgColor="0F3D6E")
BLUE = Font(name=F, color="0000FF")
START, END = dt.date(2025, 1, 1), dt.date(2026, 9, 30)

wb = Workbook()

def sheet(name, headers, rows, widths, tname, fmts=None):
    ws = wb.create_sheet(name)
    ws.append(headers)
    for r in rows: ws.append(r)
    for c, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(c)].width = w
    for cell in ws[1]:
        cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    n = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for cell in row: cell.font = Font(name=F)
    if fmts:
        for col, f in fmts.items():
            for (cell,) in ws.iter_rows(min_row=2, max_row=n, min_col=col, max_col=col): cell.number_format = f
    t = Table(displayName=tname, ref=f"A1:{get_column_letter(len(headers))}{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A2"
    return ws

# ---------- dProdutos ----------
CATS = {
 "Moda Feminina": (["Vestido Midi","Blusa Cropped","Calça Wide Leg","Saia Plissada","Macacão Longo","Blazer Alfaiataria","Cardigan Tricô","Short Jeans"], (59, 249)),
 "Moda Masculina": (["Camiseta Básica","Camisa Social","Bermuda Sarja","Calça Chino","Jaqueta Jeans","Polo Piquet","Moletom Canguru"], (49, 269)),
 "Calçados": (["Tênis Casual","Sandália Rasteira","Bota Coturno","Sapatênis","Chinelo Slide","Scarpin"], (69, 329)),
 "Acessórios": (["Bolsa Transversal","Relógio Digital","Óculos de Sol","Cinto Couro","Mochila Urbana","Boné Aba Curva","Carteira Slim"], (29, 299)),
 "Casa": (["Jogo de Cama Casal","Kit Toalhas","Luminária LED","Organizador Multiuso","Tapete Sala","Almofada Decorativa"], (39, 259)),
 "Eletrônicos": (["Fone Bluetooth","Smartwatch","Caixa de Som Portátil","Carregador Turbo","Power Bank 10000mAh","Suporte Celular"], (39, 399)),
 "Beleza": (["Kit Skincare","Perfume 100ml","Secador de Cabelo","Prancha Alisadora","Paleta de Sombras","Hidratante Corporal"], (29, 289)),
}
FORN = ["Dropi Distribuidora","Atacado Brás SP","Fornecedor Goiânia","Import Sul","Fábrica Própria"]
prods = []
pid = 1
for cat, (names, (lo, hi)) in CATS.items():
    for nm in names:
        price = round(random.uniform(lo, hi)) - 0.10
        cost = round(price * random.uniform(0.38, 0.62), 2)
        prods.append(dict(id=f"P{pid:03d}", sku=f"SKU-{cat[:3].upper()}-{pid:03d}", nome=nm, cat=cat,
                          forn=random.choice(FORN), custo=cost, preco=price,
                          pop=random.uniform(0.3, 3.0), emin=random.choice([10, 15, 20, 25, 30])))
        pid += 1

# ---------- dVendedores ----------
vends = [("V01","Ana Souza","Marketplace","Sudeste"),("V02","Bruno Lima","Loja Virtual","Sudeste"),
         ("V03","Carla Mendes","Loja Virtual","Sul"),("V04","Diego Rocha","WhatsApp","Nordeste"),
         ("V05","Elisa Martins","Instagram","Centro-Oeste"),("V06","Felipe Alves","Marketplace","Sul"),
         ("V07","Gabriela Nunes","WhatsApp","Norte"),("V08","Henrique Costa","Instagram","Nordeste")]
vw = [1.4,1.6,1.0,0.8,1.1,1.2,0.6,0.9]

# ---------- dClientes ----------
NOMES = ["Ana","Beatriz","Camila","Daniela","Eduarda","Fernanda","Gabriela","Helena","Isabela","Juliana","Larissa","Mariana","Natália","Patrícia","Rafaela","Sofia","Thaís","Vanessa",
         "André","Bruno","Carlos","Daniel","Eduardo","Felipe","Gustavo","Henrique","Igor","João","Lucas","Marcos","Pedro","Rafael","Thiago","Vinícius"]
SOBR = ["Silva","Santos","Oliveira","Souza","Rodrigues","Ferreira","Alves","Pereira","Lima","Gomes","Costa","Ribeiro","Martins","Carvalho","Almeida","Lopes","Barbosa","Rocha","Dias","Moreira"]
CID = [("São Paulo","SP","Sudeste"),("Campinas","SP","Sudeste"),("Rio de Janeiro","RJ","Sudeste"),("Belo Horizonte","MG","Sudeste"),
       ("Curitiba","PR","Sul"),("Porto Alegre","RS","Sul"),("Florianópolis","SC","Sul"),("Salvador","BA","Nordeste"),
       ("Recife","PE","Nordeste"),("Fortaleza","CE","Nordeste"),("Goiânia","GO","Centro-Oeste"),("Brasília","DF","Centro-Oeste"),
       ("Manaus","AM","Norte"),("Belém","PA","Norte")]
CW = [8,3,5,4,3,3,2,3,2,2,2,3,1,1]
clis = []
for i in range(1, 451):
    n, s = random.choice(NOMES), random.choice(SOBR)
    c = random.choices(CID, CW)[0]
    cad = START + dt.timedelta(days=random.randint(-400, 500))
    clis.append(dict(id=f"C{i:04d}", nome=f"{n} {s}", email=f"{n.lower()}.{s.lower()}{i}@exemplo.com".replace("í","i").replace("é","e").replace("ú","u").replace("ã","a").replace("á","a"),
                     cid=c[0], uf=c[1], reg=c[2], cad=cad, gen="F" if NOMES.index(n) < 18 else "M"))

# ---------- fVendas ----------
SEAS = {1:0.85,2:0.8,3:0.95,4:0.95,5:1.15,6:1.0,7:0.95,8:1.05,9:1.0,10:1.05,11:1.7,12:1.55}
CANAIS = ["Loja Virtual (Nuvemshop)","Loja Virtual (Yampi)","Mercado Livre","Shopee","Instagram","WhatsApp"]
CANW = [30,15,20,18,9,8]
PAGTO = ["Pix","Cartão de Crédito","Boleto","Cartão de Débito"]
PAGW = [45,40,8,7]
vendas = []
sold = defaultdict(int)
ped = 1
d = START
while d <= END:
    growth = 1 + (d - START).days / 365 * 0.25
    n_ped = max(0, int(random.gauss(7 * SEAS[d.month] * growth * (1.25 if d.weekday() in (0, 6) else 1), 2)))
    if d.month == 11 and 24 <= d.day <= 30: n_ped *= 3  # Black Friday
    for _ in range(n_ped):
        cli = random.choice(clis)
        v = random.choices(vends, vw)[0]
        canal = random.choices(CANAIS, CANW)[0]
        pag = random.choices(PAGTO, PAGW)[0]
        st = random.choices(["Entregue","Cancelado","Devolvido"], [91, 6, 3])[0]
        if (END - d).days < 7 and st == "Entregue": st = "Em trânsito"
        frete = 0 if random.random() < 0.35 else round(random.uniform(9.9, 34.9), 2)
        items = random.choices(prods, [p["pop"] for p in prods], k=random.choices([1,2,3,4],[60,25,10,5])[0])
        for it_no, p in enumerate(items, 1):
            q = random.choices([1,2,3],[78,17,5])[0]
            desc = random.choices([0,0.05,0.10,0.15,0.20],[55,15,15,10,5])[0]
            if d.month == 11 and 24 <= d.day <= 30: desc = max(desc, 0.20)
            vendas.append([f"PED{ped:06d}", it_no, d, cli["id"], p["id"], v[0], canal, pag, st, q, p["preco"], desc,
                           frete if it_no == 1 else 0])
            if st != "Cancelado": sold[p["id"]] += q
        ped += 1
    d += dt.timedelta(days=1)

# ---------- fCompras (entradas de estoque) ----------
compras = []
init = {}
cid_ = 1
for p in prods:
    init[p["id"]] = random.randint(40, 120)
    target_final = random.choice([random.randint(0, p["emin"]-1)] * 1 + [random.randint(p["emin"], p["emin"]*4)] * 3)
    need = sold[p["id"]] + target_final - init[p["id"]]
    m = START
    while need > 0:
        lot = min(need, random.choice([30, 50, 80, 100, 150]))
        dc = m + dt.timedelta(days=random.randint(0, 25))
        if dc > END: dc = END - dt.timedelta(days=random.randint(1, 20))
        compras.append([f"NF{cid_:05d}", dc, p["id"], p["forn"], lot, round(p["custo"] * random.uniform(0.95, 1.05), 2)])
        need -= lot; cid_ += 1
        m += dt.timedelta(days=random.randint(20, 60))
        if m > END: m = START + dt.timedelta(days=random.randint(0, 600))
compras.sort(key=lambda r: r[1])
for i, r in enumerate(compras, 1): r[0] = f"NF{i:05d}"

# ======== Escreve as abas ========
NV = len(vendas) + 1
NP = len(prods) + 1
NC = len(compras) + 1
MONEY = 'R$ #,##0.00'

ws = sheet("dProdutos", ["ID_Produto","SKU","Produto","Categoria","Fornecedor","Custo_Unit","Preco_Venda","Margem_%","Estoque_Minimo"],
      [[p["id"],p["sku"],p["nome"],p["cat"],p["forn"],p["custo"],p["preco"],f"=(G{i}-F{i})/G{i}",p["emin"]] for i,p in enumerate(prods,2)],
      [11,16,24,16,20,12,12,10,12], "tbProdutos", {6:MONEY,7:MONEY,8:"0.0%"})
for i in range(2, NP+1):
    for c in "FGI": ws[f"{c}{i}"].font = BLUE

sheet("dClientes", ["ID_Cliente","Nome","Email","Genero","Cidade","UF","Regiao","Data_Cadastro"],
      [[c["id"],c["nome"],c["email"],c["gen"],c["cid"],c["uf"],c["reg"],c["cad"]] for c in clis],
      [11,24,34,8,18,6,14,14], "tbClientes", {8:"DD/MM/YYYY"})

sheet("dVendedores", ["ID_Vendedor","Vendedor","Canal_Principal","Regiao","Meta_Mensal"],
      [[*v, m] for v, m in zip(vends, [28000,32000,20000,15000,22000,24000,12000,18000])],
      [12,20,18,14,14], "tbVendedores", {5:MONEY})

days = (END - START).days + 1
MESES = '{"Jan","Fev","Mar","Abr","Mai","Jun","Jul","Ago","Set","Out","Nov","Dez"}'
sheet("dCalendario", ["Data","Ano","Mes_Num","Mes_Nome","Ano_Mes","Trimestre","Dia_Semana","Fim_de_Semana"],
      [[START + dt.timedelta(days=k), f"=YEAR(A{k+2})", f"=MONTH(A{k+2})", f"=INDEX({MESES},C{k+2})",
        f'=TEXT(A{k+2},"yyyy-mm")', f'="T"&ROUNDUP(C{k+2}/3,0)', f'=INDEX({{"Dom","Seg","Ter","Qua","Qui","Sex","Sáb"}},WEEKDAY(A{k+2}))',
        f'=IF(OR(WEEKDAY(A{k+2})=1,WEEKDAY(A{k+2})=7),"Sim","Não")'] for k in range(days)],
      [12,8,9,10,10,10,11,13], "tbCalendario", {1:"DD/MM/YYYY"})

rows = []
for i, r in enumerate(vendas, 2):
    rows.append(r + [f"=J{i}*K{i}", f"=N{i}*L{i}", f"=N{i}-O{i}",
                     f"=J{i}*INDEX(dProdutos!$F$2:$F${NP},MATCH(E{i},dProdutos!$A$2:$A${NP},0))",
                     f'=IF(I{i}="Cancelado",0,P{i}-Q{i})',
                     f"=INDEX(dProdutos!$D$2:$D${NP},MATCH(E{i},dProdutos!$A$2:$A${NP},0))"])
sheet("fVendas", ["ID_Pedido","Item","Data","ID_Cliente","ID_Produto","ID_Vendedor","Canal","Forma_Pagamento","Status",
                  "Qtd","Preco_Unit","Desconto_%","Frete","Valor_Bruto","Valor_Desconto","Valor_Liquido","Custo_Total","Lucro","Categoria"],
      rows, [12,6,11,11,11,11,24,18,12,6,11,10,9,12,12,12,12,11,16], "tbVendas",
      {3:"DD/MM/YYYY",11:MONEY,12:"0%",13:MONEY,14:MONEY,15:MONEY,16:MONEY,17:MONEY,18:MONEY})

sheet("fCompras", ["NF_Entrada","Data_Entrada","ID_Produto","Fornecedor","Qtd","Custo_Unit","Custo_Total"],
      [r + [f"=E{i}*F{i}"] for i, r in enumerate(compras, 2)], [12,13,11,20,7,11,12], "tbCompras",
      {2:"DD/MM/YYYY",6:MONEY,7:MONEY})

est = []
for i, p in enumerate(prods, 2):
    est.append([p["id"], f"=INDEX(dProdutos!$C$2:$C${NP},MATCH(A{i},dProdutos!$A$2:$A${NP},0))", init[p["id"]],
                f"=SUMIFS(fCompras!$E$2:$E${NC},fCompras!$C$2:$C${NC},A{i})",
                f'=SUMIFS(fVendas!$J$2:$J${NV},fVendas!$E$2:$E${NV},A{i},fVendas!$I$2:$I${NV},"<>Cancelado")',
                f"=C{i}+D{i}-E{i}",
                f"=INDEX(dProdutos!$I$2:$I${NP},MATCH(A{i},dProdutos!$A$2:$A${NP},0))",
                f'=IF(F{i}<=0,"SEM ESTOQUE",IF(F{i}<G{i},"REPOR","OK"))',
                f"=F{i}*INDEX(dProdutos!$F$2:$F${NP},MATCH(A{i},dProdutos!$A$2:$A${NP},0))"])
ws = sheet("fEstoque", ["ID_Produto","Produto","Estoque_Inicial","Entradas","Saidas","Estoque_Atual","Estoque_Minimo","Status","Valor_em_Estoque"],
      est, [11,24,13,10,10,13,13,13,15], "tbEstoque", {9:MONEY})
for i in range(2, NP+1): ws[f"C{i}"].font = BLUE
rng = f"H2:H{NP}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"REPOR"'], fill=PatternFill("solid", fgColor="FFE699")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"SEM ESTOQUE"'], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"OK"'], fill=PatternFill("solid", fgColor="C6EFCE")))

# ---------- Base_Suja (treino do Pacote 1) ----------
dirty = []
for c in random.sample(clis, 220):
    nome = c["nome"]
    nome = random.choice([nome, nome.upper(), nome.lower(), "  " + nome + " ", nome.replace(" ", "  ")])
    cid = random.choice([c["cid"], c["cid"].upper(), c["cid"].lower(), c["cid"].replace("ã","a").replace("é","e").replace("ó","o")])
    uf = random.choice([c["uf"], c["uf"].lower(), {"SP":"São Paulo","RJ":"Rio","MG":"Minas"}.get(c["uf"], c["uf"])])
    cad = random.choice([c["cad"].strftime("%d/%m/%Y"), c["cad"].strftime("%Y-%m-%d"), c["cad"].strftime("%d-%m-%y"), c["cad"].strftime("%d.%m.%Y")])
    tel = f"{random.randint(11,99)}9{random.randint(1000,9999)}{random.randint(1000,9999)}"
    tel = random.choice([tel, f"({tel[:2]}) {tel[2:7]}-{tel[7:]}", f"{tel[:2]} {tel[2:]}", f"+55{tel}"])
    val = round(random.uniform(50, 3000), 2)
    val = random.choice([val, f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", "."), str(val).replace(".", ","), ""])
    dirty.append([c["id"], nome, c["email"].upper() if random.random() < 0.2 else c["email"], tel, cid, uf, cad, val])
dirty += random.sample(dirty, 25)
dirty += [[None]*8 for _ in range(6)]
random.shuffle(dirty)
ws = wb.create_sheet("Base_Suja")
ws.append(["cod cliente","NOME CLIENTE","e-mail","Telefone","cidade","Estado","dt cadastro","Total Gasto"])
for r in dirty: ws.append(r)
for c, w in enumerate([12,28,34,18,18,12,14,14], 1): ws.column_dimensions[get_column_letter(c)].width = w
for row in ws.iter_rows():
    for cell in row: cell.font = Font(name=F)

# ---------- Resumo ----------
ws = wb.create_sheet("Resumo", 1)
ws.column_dimensions["A"].width = 30
for c in "BCDEF": ws.column_dimensions[c].width = 16
ws["A1"] = "Resumo Gerencial – Base Fictícia de E-commerce"; ws["A1"].font = Font(name=F, bold=True, size=14, color="0F3D6E")
ws["A2"] = "Valores calculados por fórmula a partir de fVendas (exclui pedidos cancelados)."; ws["A2"].font = Font(name=F, italic=True, color="666666")
V = lambda col: f"fVendas!${col}$2:${col}${NV}"
kpis = [("Faturamento Líquido", f'=SUMIFS({V("P")},{V("I")},"<>Cancelado")', MONEY),
        ("Lucro Bruto", f'=SUM({V("R")})', MONEY),
        ("Margem %", "=IF(B4=0,0,B5/B4)", "0.0%"),
        ("Nº de Pedidos", f'=COUNTIFS({V("B")},1,{V("I")},"<>Cancelado")', "#,##0"),
        ("Ticket Médio", "=IF(B7=0,0,B4/B7)", MONEY),
        ("Itens Vendidos", f'=SUMIFS({V("J")},{V("I")},"<>Cancelado")', "#,##0"),
        ("Produtos para Repor", f'=COUNTIF(fEstoque!$H$2:$H${NP},"REPOR")+COUNTIF(fEstoque!$H$2:$H${NP},"SEM ESTOQUE")', "0")]
for k, (lab, f, fmt) in enumerate(kpis, 4):
    ws[f"A{k}"] = lab; ws[f"B{k}"] = f; ws[f"B{k}"].number_format = fmt
    ws[f"A{k}"].font = Font(name=F, bold=True); ws[f"B{k}"].font = Font(name=F, bold=True, color="0F3D6E", size=12)

def hdr(r, labels):
    for c, l in enumerate(labels, 1):
        cell = ws.cell(r, c, l); cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR

r0 = 13
hdr(r0, ["Mês","Faturamento","Lucro","Margem %"])
m = dt.date(2025, 1, 1); r = r0 + 1
while m <= END:
    nxt = (m.replace(day=28) + dt.timedelta(days=4)).replace(day=1)
    ws[f"A{r}"] = m; ws[f"A{r}"].number_format = "mmm/yyyy"
    ws[f"B{r}"] = f'=SUMIFS({V("P")},{V("C")},">="&A{r},{V("C")},"<"&EDATE(A{r},1),{V("I")},"<>Cancelado")'
    ws[f"C{r}"] = f'=SUMIFS({V("R")},{V("C")},">="&A{r},{V("C")},"<"&EDATE(A{r},1))'
    ws[f"D{r}"] = f"=IF(B{r}=0,0,C{r}/B{r})"
    ws[f"B{r}"].number_format = ws[f"C{r}"].number_format = MONEY; ws[f"D{r}"].number_format = "0.0%"
    for c in "ABCD": ws[f"{c}{r}"].font = Font(name=F)
    m = nxt; r += 1
r_end_m = r - 1
ch = LineChart(); ch.title = "Faturamento Mensal"; ch.height, ch.width = 7.5, 16; ch.legend = None
ch.add_data(Reference(ws, min_col=2, min_row=r0, max_row=r_end_m), titles_from_data=True)
ch.set_categories(Reference(ws, min_col=1, min_row=r0+1, max_row=r_end_m))
ws.add_chart(ch, "F13")

r1 = r_end_m + 3
hdr(r1, ["Categoria","Faturamento","Lucro","Margem %"])
for k, cat in enumerate(CATS, r1 + 1):
    ws[f"A{k}"] = cat
    ws[f"B{k}"] = f'=SUMIFS({V("P")},{V("S")},A{k},{V("I")},"<>Cancelado")'
    ws[f"C{k}"] = f'=SUMIFS({V("R")},{V("S")},A{k})'
    ws[f"D{k}"] = f"=IF(B{k}=0,0,C{k}/B{k})"
    ws[f"B{k}"].number_format = ws[f"C{k}"].number_format = MONEY; ws[f"D{k}"].number_format = "0.0%"
    for c in "ABCD": ws[f"{c}{k}"].font = Font(name=F)
r1e = r1 + len(CATS)
bc = BarChart(); bc.type = "bar"; bc.title = "Faturamento por Categoria"; bc.height, bc.width = 7.5, 16; bc.legend = None
bc.add_data(Reference(ws, min_col=2, min_row=r1, max_row=r1e), titles_from_data=True)
bc.set_categories(Reference(ws, min_col=1, min_row=r1+1, max_row=r1e))
ws.add_chart(bc, f"F{r1}")

r2 = r1e + 3
hdr(r2, ["Canal","Faturamento","Nº Pedidos","Ticket Médio"])
for k, can in enumerate(CANAIS, r2 + 1):
    ws[f"A{k}"] = can
    ws[f"B{k}"] = f'=SUMIFS({V("P")},{V("G")},A{k},{V("I")},"<>Cancelado")'
    ws[f"C{k}"] = f'=COUNTIFS({V("G")},A{k},{V("B")},1,{V("I")},"<>Cancelado")'
    ws[f"D{k}"] = f"=IF(C{k}=0,0,B{k}/C{k})"
    ws[f"B{k}"].number_format = ws[f"D{k}"].number_format = MONEY
    for c in "ABCD": ws[f"{c}{k}"].font = Font(name=F)

# ---------- LEIA-ME ----------
ws = wb["Sheet"]; ws.title = "LEIA-ME"
ws.column_dimensions["A"].width = 22; ws.column_dimensions["B"].width = 95
lines = [("Base Fictícia de E-commerce – Estudo e Portfólio", None),
 (f"Período: {START:%d/%m/%Y} a {END:%d/%m/%Y} · {len(vendas):,} linhas de venda · {ped-1:,} pedidos · {len(prods)} produtos · {len(clis)} clientes".replace(",", "."), None),
 ("Todos os dados são FICTÍCIOS (gerados aleatoriamente). Podem ser usados livremente em portfólio.", None),
 ("", None), ("ABA", "CONTEÚDO / USO"),
 ("Resumo", "KPIs, faturamento mensal, por categoria e por canal (tudo em fórmulas) + gráficos. Modelo para o Pacote 2."),
 ("dProdutos", "Dimensão de produtos: categoria, fornecedor, custo, preço, margem e estoque mínimo."),
 ("dClientes", "Dimensão de clientes: cidade, UF, região, data de cadastro."),
 ("dVendedores", "Dimensão de vendedores com canal principal e meta mensal (para medidas de % da meta no Power BI)."),
 ("dCalendario", "Tabela calendário pronta (ano, mês, trimestre, dia da semana) – use no Power BI."),
 ("fVendas", "Tabela fato: 1 linha por item de pedido. Item = 1 marca a 1ª linha do pedido (use para contar pedidos)."),
 ("fCompras", "Entradas de estoque (notas fiscais de compra dos fornecedores)."),
 ("fEstoque", "Posição de estoque: inicial + entradas − saídas, com alerta OK / REPOR / SEM ESTOQUE."),
 ("Base_Suja", "Cadastro BAGUNÇADO de propósito (duplicados, linhas vazias, datas/telefones/valores fora do padrão). Treino do Pacote 1 com Power Query."),
 ("", None), ("CONVENÇÕES", None),
 ("Texto azul", "Valores digitados (premissas) que você pode alterar: custo, preço, estoque mínimo, estoque inicial."),
 ("Texto preto", "Fórmulas – recalculam sozinhas."),
 ("Status", "Pedidos 'Cancelado' não entram no faturamento, lucro nem nas saídas de estoque. 'Devolvido' foi mantido como venda (exercício: tratar devoluções)."),
 ("Frete", "Cobrado do cliente, registrado apenas na 1ª linha (Item = 1) de cada pedido; não entra no Valor_Liquido."),
 ("Sazonalidade", "Pico em novembro (Black Friday, 24 a 30/11 com volume ×3 e desconto mínimo de 20%) e dezembro; crescimento de ~25% ao ano."),
 ("", None), ("POWER BI", "Obter Dados > Excel > marque as tabelas tb* (tbVendas, tbProdutos, tbClientes, tbVendedores, tbCalendario, tbCompras)."),
 ("", "Relacione: tbVendas[ID_Produto]→tbProdutos, [ID_Cliente]→tbClientes, [ID_Vendedor]→tbVendedores, [Data]→tbCalendario[Data]."),
]
for i, (a, b) in enumerate(lines, 1):
    ws[f"A{i}"] = a; ws[f"A{i}"].font = Font(name=F, bold=True)
    if b: ws[f"B{i}"] = b; ws[f"B{i}"].font = Font(name=F); ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(name=F, bold=True, size=14, color="0F3D6E")
for i in (2, 3): ws.merge_cells(f"A{i}:B{i}"); ws[f"A{i}"].font = Font(name=F, italic=True)
for c in "AB": ws[f"{c}5"].font = Font(name=F, bold=True, color="FFFFFF"); ws[f"{c}5"].fill = HDR
ws["A16"].font = Font(name=F, bold=True, color="0F3D6E")
ws["B17"].font = Font(name=F, color="0000FF")

wb.save(OUT)
print(len(vendas), ped-1, len(compras))
