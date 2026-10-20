# Gera a planilha de acompanhamento de estudos (Acompanhamento_Estudos.xlsx) na mesma pasta deste script.
# Requisito: pip install openpyxl  |  Uso: python gerar_planilha_estudos.py
import os, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "Acompanhamento_Estudos.xlsx")
F = "Arial"
NAVY = "0F3D6E"
HDR = PatternFill("solid", fgColor=NAVY)
INPUT = PatternFill("solid", fgColor="FFF2CC")
BLUE = Font(name=F, color="0000FF")
thin = Side(style="thin", color="D6DDE6")

# (Bloco, Pacotes, Tópico, Nível, Semana, Fonte sugerida, Exercício prático)
T = [
 ("Excel", "1", "Tabelas formatadas (Ctrl+T) e boas práticas de nomes de colunas", "Básico", 1, "Microsoft Learn / Hashtag", "Converter dClientes em tabela e renomear colunas"),
 ("Excel", "1", "Remover duplicadas (escolhendo as colunas-chave)", "Básico", 1, "Hashtag Treinamentos", "Base_Suja: remover clientes repetidos"),
 ("Excel", "1", "Funções de texto: ARRUMAR, PRI.MAIÚSCULA, SUBSTITUIR, ESQUERDA/DIREITA/EXT.TEXTO", "Básico", 1, "Hashtag Treinamentos", "Base_Suja: padronizar NOME CLIENTE e cidade"),
 ("Excel", "1", "Texto para Colunas, TEXTOANTES/TEXTODEPOIS, concatenar", "Básico", 1, "Hashtag Treinamentos", "Separar nome e sobrenome em dClientes"),
 ("Excel", "1,2", "Funções de data: HOJE, DATA, MÊS, ANO, FIMMÊS, DATADIF, DIATRABALHOTOTAL", "Básico", 1, "Microsoft Learn", "Calcular tempo de cadastro dos clientes"),
 ("Excel", "1,2", "Funções lógicas: SE, SES, E, OU, SEERRO", "Básico", 1, "Hashtag Treinamentos", "Classificar pedidos por faixa de valor"),
 ("Excel", "2", "Referências relativas e absolutas ($)", "Básico", 1, "Hashtag Treinamentos", "Montar tabela de faturamento por mês"),
 ("Excel", "2", "PROCX e ÍNDICE + CORRESP", "Intermediário", 2, "Hashtag Treinamentos", "Trazer a Categoria do produto para fVendas"),
 ("Excel", "2", "SOMASES, CONT.SES, MÉDIASES com critérios de data", "Intermediário", 2, "Hashtag Treinamentos", "Refazer os KPIs da aba Resumo"),
 ("Excel", "1,2", "Tabela dinâmica: linhas, colunas, valores, filtros", "Básico", 2, "Microsoft Learn", "Faturamento por Categoria x Canal"),
 ("Excel", "2", "Gráfico dinâmico, Segmentação de Dados e Linha do Tempo", "Intermediário", 2, "Karine Lago", "Gráfico mensal com segmentação por canal"),
 ("Excel", "2", "Formatação condicional (regras, barras de dados, ícones)", "Básico", 2, "Hashtag Treinamentos", "Alerta de estoque na aba fEstoque"),
 ("Excel", "2", "Validação de dados (listas suspensas) e Proteção de planilha", "Básico", 2, "Microsoft Learn", "Lista de Status nesta planilha / proteger fórmulas"),
 ("Excel", "2", "Matrizes dinâmicas: FILTRO, ÚNICO, CLASSIFICAR, LET (Microsoft 365)", "Intermediário", 2, "Microsoft Learn", "Top 10 produtos com CLASSIFICAR + FILTRO"),
 ("Power Query", "1,2,3", "Importar dados: Excel, CSV, Web, Pasta", "Básico", 3, "Microsoft Learn (PL-300)", "Importar a aba Base_Suja"),
 ("Power Query", "1,2,3", "Tipos de dados e Alterar Tipo com Localidade", "Básico", 3, "Microsoft Learn (PL-300)", "Base_Suja: corrigir datas em 4 formatos"),
 ("Power Query", "1,2,3", "Limpar, Aparar, Substituir valores, Remover linhas em branco/erros", "Básico", 3, "Karine Lago", "Base_Suja: limpar telefone e Total Gasto"),
 ("Power Query", "1,2,3", "Dividir e mesclar colunas, Extrair, Coluna condicional, Coluna personalizada", "Intermediário", 3, "Leonardo Karpinski", "Criar coluna Faixa de Gasto"),
 ("Power Query", "1,3", "Mesclar consultas (JOIN) e Acrescentar consultas (UNION)", "Intermediário", 3, "Leonardo Karpinski", "Juntar fVendas com dProdutos"),
 ("Power Query", "1,3", "Transformar colunas em linhas (unpivot) e Dinamizar", "Intermediário", 3, "Karine Lago", "Despivotar a tabela mensal da aba Resumo"),
 ("Power Query", "1,4", "Obter dados de Pasta (vários arquivos) e consulta reaproveitável", "Intermediário", 3, "Microsoft Learn (PL-300)", "Separar fVendas por ano em CSVs e combinar"),
 ("Power Query", "1,3", "Linguagem M: editor avançado e parâmetros", "Avançado", 3, "Microsoft Learn", "Parametrizar o caminho da pasta de origem"),
 ("Negócio", "2", "KPIs de e-commerce: faturamento, ticket médio, margem", "Básico", 4, "Blogs Nuvemshop / Yampi", "Conferir os KPIs da aba Resumo"),
 ("Negócio", "2", "Curva ABC de produtos", "Intermediário", 4, "Hashtag Treinamentos", "Montar curva ABC com fVendas"),
 ("Negócio", "2", "Giro de estoque, estoque mínimo e ponto de pedido", "Intermediário", 4, "Blogs de gestão de estoque", "Calcular giro em fEstoque"),
 ("Excel", "2", "Montagem de painel (dashboard) no Excel: layout, cartões de KPI, gráficos", "Intermediário", 4, "Karine Lago", "AMOSTRA 2 do portfólio: painel de vendas e estoque"),
 ("Dados", "2,3", "Exportações das plataformas (Nuvemshop, Yampi, Shopee, Mercado Livre, Bling)", "Intermediário", 4, "Central de ajuda de cada plataforma", "Mapear colunas de um relatório real"),
 ("Power BI", "3", "Interface do Power BI Desktop e Obter Dados", "Básico", 5, "Microsoft Learn (PL-300)", "Importar as tabelas tb* da base"),
 ("Power BI", "3", "Esquema estrela: tabelas fato x dimensão", "Avançado", 5, "Microsoft Learn (PL-300)", "Desenhar o modelo da base de e-commerce"),
 ("Power BI", "3", "Relacionamentos: cardinalidade 1:N e direção do filtro", "Avançado", 5, "Leonardo Karpinski", "Relacionar tbVendas com as 4 dimensões"),
 ("Power BI", "3", "Tabela calendário (marcar como tabela de datas)", "Intermediário", 5, "Leonardo Karpinski", "Usar tbCalendario no modelo"),
 ("Power BI", "3", "Visuais: cartões, barras, linhas, matriz, mapa", "Básico", 5, "Microsoft Learn (PL-300)", "Página 1 do dashboard"),
 ("Power BI", "3", "Segmentações, interação entre visuais e sincronização entre páginas", "Básico", 5, "Karine Lago", "Filtros por mês, vendedor e categoria"),
 ("Power BI", "3", "DAX básico: SUM, COUNTROWS, DISTINCTCOUNT, DIVIDE, AVERAGE", "Intermediário", 6, "Microsoft Learn (PL-300)", "Medidas Faturamento, Pedidos, Ticket Médio"),
 ("Power BI", "3", "Medida x coluna calculada; contexto de linha x contexto de filtro", "Avançado", 6, "SQLBI (sqlbi.com)", "Explicar por escrito a diferença"),
 ("Power BI", "3", "CALCULATE, FILTER, ALL, ALLEXCEPT, REMOVEFILTERS", "Avançado", 6, "SQLBI / Leonardo Karpinski", "% de participação por categoria"),
 ("Power BI", "3", "Iteradores: SUMX, AVERAGEX (margem e lucro)", "Intermediário", 6, "SQLBI", "Medida Lucro = SUMX(Qtd × (Preço − Custo))"),
 ("Power BI", "3", "Inteligência de tempo: TOTALYTD, SAMEPERIODLASTYEAR, DATEADD", "Avançado", 7, "SQLBI / Microsoft Learn", "Faturamento x mesmo mês do ano anterior"),
 ("Power BI", "3", "RANKX, TOPN, SWITCH, VAR/RETURN", "Avançado", 7, "SQLBI", "Ranking de vendedores e % da meta"),
 ("Power BI", "3", "Tema JSON, identidade visual e layout profissional", "Intermediário", 7, "Karine Lago", "Tema com as cores de uma marca fictícia"),
 ("Power BI", "3", "Drill-through, dicas de ferramenta e indicadores (bookmarks)", "Intermediário", 7, "Microsoft Learn (PL-300)", "Página de detalhe do produto"),
 ("Power BI", "3", "Montar o dashboard executivo completo (3 páginas)", "Intermediário", 7, "—", "AMOSTRA 3 do portfólio"),
 ("Power BI", "3,4", "Power BI Service: workspace, publicar e compartilhar", "Intermediário", 8, "Microsoft Learn (PL-300)", "Publicar o dashboard (conta gratuita/teste Pro)"),
 ("Power BI", "3,4", "Atualização agendada e Gateway de dados local", "Avançado", 8, "Microsoft Learn", "Agendar atualização diária"),
 ("Power BI", "3", "Segurança por usuário (RLS) e USERPRINCIPALNAME", "Avançado", 8, "Microsoft Learn", "RLS por região dos vendedores"),
 ("Power BI", "3", "Licenças: Gratuita, Pro, Premium/Fabric", "Básico", 8, "Microsoft Learn", "Resumo de 5 linhas para explicar ao cliente"),
 ("Dados", "4", "Checklist de conferência de dados (totais, duplicados, valores fora do padrão)", "Intermediário", 8, "—", "Criar o checklist do Pacote 4"),
 ("Dados", "4", "Diagnóstico de falhas de atualização (credenciais, colunas renomeadas)", "Avançado", 8, "Microsoft Learn", "Renomear uma coluna na fonte e corrigir o erro"),
 ("Dados", "3,E", "SQL: SELECT, JOIN, GROUP BY, views", "Intermediário", 6, "SQLBolt / Mode SQL Tutorial", "Importar a base em SQLite/MySQL e consultar"),
 ("Dados", "3", "Conectar Power BI a MySQL, PostgreSQL e SQL Server", "Intermediário", 8, "Microsoft Learn", "Conectar o Power BI ao banco criado"),
 ("Dados", "3", "Google Sheets como fonte de dados", "Básico", 5, "Microsoft Learn", "Conectar uma planilha do Google"),
 ("Python", "E", "pandas: ler, limpar, juntar e agrupar dados", "Intermediário", 4, "Documentação pandas / Kaggle Learn", "Limpar a Base_Suja com pandas"),
 ("Python", "E", "openpyxl: gerar Excel formatado com fórmulas", "Intermediário", 6, "Documentação openpyxl", "Estudar o script gerar_base.py"),
 ("Python", "E", "requests: consumir APIs de e-commerce", "Intermediário", 7, "Documentação requests / API Nuvemshop", "Ler a documentação da API da Nuvemshop"),
 ("Python", "E", "Agendar scripts no Agendador de Tarefas do Windows", "Básico", 8, "Microsoft Docs", "Agendar o gerar_base.py"),
 ("Negócio", "Todos", "Levantamento de requisitos com o cliente", "Básico", 4, "—", "Roteiro de 6 perguntas (seção 5 do portfólio)"),
 ("Negócio", "Todos", "Storytelling com dados", "Intermediário", 7, "Livro: Storytelling com Dados (Cole Knaflic)", "Revisar o dashboard com menos gráficos"),
 ("Negócio", "Todos", "Gravar vídeo tutorial de entrega (OBS / Loom)", "Básico", 4, "YouTube", "Vídeo de 30s da amostra 2"),
 ("Negócio", "Todos", "LGPD: tratamento e descarte de dados do cliente", "Básico", 8, "gov.br/anpd", "Texto de política de sigilo"),
 ("Negócio", "Todos", "Precificação, escopo e revisões", "Básico", 8, "—", "Revisar preços do portfólio"),
 ("Excel", "1", "AMOSTRA 1 do portfólio: Antes x Depois da Base_Suja", "Intermediário", 3, "—", "Prints lado a lado + GIF"),
]
T.sort(key=lambda t: (t[4], ["Excel","Power Query","Negócio","Dados","Power BI","Python"].index(t[0])))
N = len(T); LAST = N + 1

wb = Workbook()

def style_header(ws, row, ncol):
    for c in range(1, ncol + 1):
        cell = ws.cell(row, c)
        cell.font = Font(name=F, bold=True, color="FFFFFF"); cell.fill = HDR
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# ---------------- Checklist ----------------
ck = wb.create_sheet("Checklist")
H = ["ID","Semana","Bloco","Pacotes","Tópico","Nível","Status","Data_Inicio","Data_Conclusao","Horas_Estudadas","Fonte_Sugerida","Exercicio_Pratico","Observacoes"]
ck.append(H)
for i, (bl, pc, tp, nv, sem, fo, ex) in enumerate(T, 1):
    r = i + 1
    ck.append([f"T{i:02d}", sem, bl, pc, tp, nv, "Não iniciado", None, None,
               f'=SUMIFS(Diario!$D:$D,Diario!$B:$B,A{r})', fo, ex, None])
widths = [7,8,13,9,58,13,14,12,14,10,30,46,30]
for c, w in enumerate(widths, 1): ck.column_dimensions[get_column_letter(c)].width = w
style_header(ck, 1, len(H)); ck.row_dimensions[1].height = 30
for row in ck.iter_rows(min_row=2, max_row=LAST):
    for cell in row:
        cell.font = Font(name=F); cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (5, 11, 12, 13))
    for c in (7, 8, 9, 13):
        row[c-1].fill = INPUT; row[c-1].font = BLUE
    row[7].number_format = row[8].number_format = "DD/MM/YYYY"
    row[9].number_format = "0.0"
    row[1].alignment = row[3].alignment = Alignment(horizontal="center", vertical="top")
t = Table(displayName="tbChecklist", ref=f"A1:{get_column_letter(len(H))}{LAST}")
t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=False)
ck.add_table(t); ck.freeze_panes = "F2"
dv = DataValidation(type="list", formula1='"Não iniciado,Estudando,Concluído,Revisar"', allow_blank=False,
                    showErrorMessage=True, errorTitle="Status inválido", error="Escolha um item da lista.")
ck.add_data_validation(dv); dv.add(f"G2:G{LAST}")
dvd = DataValidation(type="date", operator="greaterThan", formula1="DATE(2020,1,1)", allow_blank=True,
                     showErrorMessage=True, error="Digite uma data válida (dd/mm/aaaa).")
ck.add_data_validation(dvd); dvd.add(f"H2:I{LAST}")
S = f"G2:G{LAST}"
for val, color in [("Concluído","C6EFCE"),("Estudando","BDD7EE"),("Revisar","FFE699"),("Não iniciado","F2F2F2")]:
    ck.conditional_formatting.add(S, CellIsRule(operator="equal", formula=[f'"{val}"'], fill=PatternFill("solid", fgColor=color)))
ck.conditional_formatting.add(f"A2:F{LAST}", FormulaRule(formula=[f'$G2="Concluído"'], font=Font(name=F, color="808080", strike=True)))
NV = f"F2:F{LAST}"
ck.conditional_formatting.add(NV, CellIsRule(operator="equal", formula=['"Avançado"'], font=Font(name=F, bold=True, color="C00000")))
ck.conditional_formatting.add(NV, CellIsRule(operator="equal", formula=['"Intermediário"'], font=Font(name=F, bold=True, color="BF8F00")))
ck.conditional_formatting.add(NV, CellIsRule(operator="equal", formula=['"Básico"'], font=Font(name=F, bold=True, color="548235")))

# ---------------- Diario ----------------
di = wb.create_sheet("Diario")
di.append(["Data","ID_Topico","Topico (automático)","Horas","O que aprendi / dúvidas"])
style_header(di, 1, 5)
di.append([dt.date(2026, 10, 20), "T01", "=IFERROR(INDEX(Checklist!$E:$E,MATCH(B2,Checklist!$A:$A,0)),\"\")", 0,
           "EXEMPLO – substitua por um registro real (horas = 0 para não contar)"])
for r in range(3, 401):
    di.cell(r, 3, f'=IF(B{r}="","",IFERROR(INDEX(Checklist!$E:$E,MATCH(B{r},Checklist!$A:$A,0)),"ID não encontrado"))')
for c, w in enumerate([12,11,58,8,60], 1): di.column_dimensions[get_column_letter(c)].width = w
for row in di.iter_rows(min_row=2, max_row=400):
    for cell in row: cell.font = Font(name=F)
    for c in (1, 2, 4, 5): row[c-1].fill = INPUT; row[c-1].font = BLUE
    row[0].number_format = "DD/MM/YYYY"; row[3].number_format = "0.0"
for c in range(1, 6): di.cell(2, c).font = Font(name=F, italic=True, color="808080")
dvi = DataValidation(type="list", formula1=f"=Checklist!$A$2:$A${LAST}", allow_blank=True)
di.add_data_validation(dvi); dvi.add("B2:B400")
dvh = DataValidation(type="decimal", operator="between", formula1="0", formula2="16", allow_blank=True,
                     showErrorMessage=True, error="Horas entre 0 e 16.")
di.add_data_validation(dvh); dvh.add("D2:D400")
di.freeze_panes = "A2"

# ---------------- Cronograma ----------------
cr = wb.create_sheet("Cronograma")
cr["A1"] = "Cronograma de 8 semanas (1 a 2 h por dia)"; cr["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
cr["A3"] = "Data de início do estudo:"; cr["A3"].font = Font(name=F, bold=True)
cr["C3"] = dt.date(2026, 10, 26); cr["C3"].fill = INPUT; cr["C3"].font = Font(name=F, bold=True, color="0000FF"); cr["C3"].number_format = "DD/MM/YYYY"
cr["D3"] = "← altere para a sua data (segunda-feira)"; cr["D3"].font = Font(name=F, italic=True, color="808080")
hdr = ["Semana","Início","Fim","Foco","Libera","Tópicos","Concluídos","% Semana"]
for c, h in enumerate(hdr, 1): cr.cell(5, c, h)
style_header(cr, 5, len(hdr))
plan = [(1,"Excel: tabelas, fórmulas básicas, texto, data, lógica","—"),
        (2,"Excel: PROCX, SOMASES, tabela dinâmica, formatação condicional","Pacotes 1 e 2 (Básico)"),
        (3,"Power Query + AMOSTRA 1","Pacote 1 completo"),
        (4,"Painel no Excel, KPIs de e-commerce, pandas + AMOSTRA 2","Pacote 2 completo"),
        (5,"Power BI: interface, modelagem estrela, visuais","Pacote 3 (Básico)"),
        (6,"DAX: básico, CALCULATE, iteradores + SQL","Pacote 3 (Padrão)"),
        (7,"DAX de tempo, ranking, tema, drill-through + AMOSTRA 3","Pacote 3 (Padrão)"),
        (8,"Power BI Service, Gateway, RLS, manutenção, LGPD","Pacote 3 Premium + Pacote 4")]
for k, (s, foco, lib) in enumerate(plan, 6):
    cr.cell(k, 1, s); cr.cell(k, 2, f"=$C$3+7*(A{k}-1)"); cr.cell(k, 3, f"=B{k}+6")
    cr.cell(k, 4, foco); cr.cell(k, 5, lib)
    cr.cell(k, 6, f"=COUNTIF(Checklist!$B$2:$B${LAST},A{k})")
    cr.cell(k, 7, f'=COUNTIFS(Checklist!$B$2:$B${LAST},A{k},Checklist!$G$2:$G${LAST},"Concluído")')
    cr.cell(k, 8, f"=IF(F{k}=0,0,G{k}/F{k})")
    for c in range(1, 9):
        cell = cr.cell(k, c); cell.font = Font(name=F); cell.border = Border(bottom=thin)
        cell.alignment = Alignment(vertical="center", wrap_text=c in (4, 5), horizontal="center" if c in (1,6,7,8) else None)
    cr.cell(k, 2).number_format = cr.cell(k, 3).number_format = "DD/MM/YYYY"; cr.cell(k, 8).number_format = "0%"
    cr.row_dimensions[k].height = 30
cr.conditional_formatting.add("H6:H13", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E75B6"))
for c, w in enumerate([9,12,12,52,28,9,11,14], 1): cr.column_dimensions[get_column_letter(c)].width = w

# ---------------- Painel ----------------
pn = wb.create_sheet("Painel", 0)
pn["A1"] = "Painel de Estudos – Especialista em Power BI e Excel"; pn["A1"].font = Font(name=F, bold=True, size=16, color=NAVY)
pn["A2"] = "Tudo nesta aba é calculado automaticamente. Atualize apenas a aba Checklist (Status) e a aba Diario (horas)."
pn["A2"].font = Font(name=F, italic=True, color="808080")
G = f"Checklist!$G$2:$G${LAST}"; B = f"Checklist!$C$2:$C${LAST}"; P = f"Checklist!$D$2:$D${LAST}"; NIV = f"Checklist!$F$2:$F${LAST}"
cards = [("Tópicos", f"=COUNTA(Checklist!$A$2:$A${LAST})", "0"),
         ("Concluídos", f'=COUNTIF({G},"Concluído")', "0"),
         ("% Concluído", "=IF(B4=0,0,D4/B4)", "0%"),
         ("Horas estudadas", "=SUM(Diario!$D$2:$D$400)", "0.0")]
for k, (lab, f, fmt) in enumerate(cards):
    col = 1 + k * 2
    lc = pn.cell(3, col, lab); vc = pn.cell(4, col, f)
    lc.font = Font(name=F, bold=True, color="FFFFFF"); lc.fill = HDR; lc.alignment = Alignment(horizontal="center")
    vc.font = Font(name=F, bold=True, size=20, color=NAVY); vc.number_format = fmt; vc.alignment = Alignment(horizontal="center")
    vc.fill = PatternFill("solid", fgColor="F3F6FA")
    pn.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + 1)
    pn.merge_cells(start_row=4, start_column=col, end_row=4, end_column=col + 1)
pn.row_dimensions[4].height = 36
# fix: card values reference merged anchors
pn["E4"] = "=IF(A4=0,0,C4/A4)"

pn["A6"] = "Status dos tópicos"; pn["A6"].font = Font(name=F, bold=True, size=12, color=NAVY)
for k, s in enumerate(["Não iniciado","Estudando","Revisar","Concluído"], 7):
    pn.cell(k, 1, s).font = Font(name=F); c = pn.cell(k, 2, f'=COUNTIF({G},A{k})'); c.font = Font(name=F, bold=True)

r0 = 12
pn.cell(r0, 1, "Progresso por bloco de conhecimento").font = Font(name=F, bold=True, size=12, color=NAVY)
for c, h in enumerate(["Bloco","Tópicos","Concluídos","% Concluído","Horas"], 1): pn.cell(r0 + 1, c, h)
style_header(pn, r0 + 1, 5)
blocos = ["Excel","Power Query","Power BI","Dados","Python","Negócio"]
for k, bl in enumerate(blocos, r0 + 2):
    pn.cell(k, 1, bl)
    pn.cell(k, 2, f"=COUNTIF({B},A{k})")
    pn.cell(k, 3, f'=COUNTIFS({B},A{k},{G},"Concluído")')
    pn.cell(k, 4, f"=IF(B{k}=0,0,C{k}/B{k})").number_format = "0%"
    pn.cell(k, 5, f"=SUMIFS(Checklist!$J$2:$J${LAST},{B},A{k})").number_format = "0.0"
    for c in range(1, 6): pn.cell(k, c).font = Font(name=F); pn.cell(k, c).border = Border(bottom=thin)
rb_end = r0 + 1 + len(blocos)
pn.conditional_formatting.add(f"D{r0+2}:D{rb_end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E75B6"))
ch = BarChart(); ch.type = "bar"; ch.title = "% concluído por bloco"; ch.legend = None; ch.height, ch.width = 7, 13
ch.add_data(Reference(pn, min_col=4, min_row=r0 + 1, max_row=rb_end), titles_from_data=True)
ch.set_categories(Reference(pn, min_col=1, min_row=r0 + 2, max_row=rb_end))
ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 1; ch.y_axis.numFmt = "0%"
pn.add_chart(ch, "J3")

r1 = rb_end + 3
pn.cell(r1, 1, "Liberação dos pacotes para venda").font = Font(name=F, bold=True, size=12, color=NAVY)
pn.cell(r1 + 1, 1, "Nível Básico liberado = todos os tópicos Básico e Intermediário do pacote concluídos. Completo = todos, inclusive os Avançados.").font = Font(name=F, italic=True, color="808080")
for c, h in enumerate(["Pacote","Tópicos","Concluídos","% Concluído","Nível Básico","Completo (Padrão/Premium)"], 1): pn.cell(r1 + 2, c, h)
style_header(pn, r1 + 2, 6)
pacs = [("1","Pacote 1 – Limpeza e Padronização"),("2","Pacote 2 – Vendas e Estoque no Excel"),
        ("3","Pacote 3 – Dashboard Power BI"),("4","Pacote 4 – Manutenção Recorrente"),("E","Extras – Python e automação")]
for k, (cod, nome) in enumerate(pacs, r1 + 3):
    crit = f'"*{cod}*"'
    pn.cell(k, 1, nome)
    pn.cell(k, 2, f'=COUNTIF({P},{crit})+COUNTIF({P},"Todos")')
    pn.cell(k, 3, f'=COUNTIFS({P},{crit},{G},"Concluído")+COUNTIFS({P},"Todos",{G},"Concluído")')
    pn.cell(k, 4, f"=IF(B{k}=0,0,C{k}/B{k})").number_format = "0%"
    nb = f'COUNTIFS({P},{crit},{NIV},"<>Avançado")+COUNTIFS({P},"Todos",{NIV},"<>Avançado")'
    nbc = f'COUNTIFS({P},{crit},{NIV},"<>Avançado",{G},"Concluído")+COUNTIFS({P},"Todos",{NIV},"<>Avançado",{G},"Concluído")'
    pn.cell(k, 5, f'=IF({nbc}={nb},"✔ PODE VENDER","Em preparação")')
    pn.cell(k, 6, f'=IF(C{k}=B{k},"✔ PODE VENDER","Em preparação")')
    for c in range(1, 7): pn.cell(k, c).font = Font(name=F); pn.cell(k, c).border = Border(bottom=thin)
rp_end = r1 + 2 + len(pacs)
rngp = f"E{r1+3}:F{rp_end}"
pn.conditional_formatting.add(rngp, CellIsRule(operator="equal", formula=['"✔ PODE VENDER"'], fill=PatternFill("solid", fgColor="C6EFCE"), font=Font(name=F, bold=True, color="006100")))
pn.conditional_formatting.add(rngp, CellIsRule(operator="equal", formula=['"Em preparação"'], fill=PatternFill("solid", fgColor="FFF2CC"), font=Font(name=F, color="7F6000")))
pn.conditional_formatting.add(f"D{r1+3}:D{rp_end}", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="2E75B6"))
for c, w in enumerate([36,11,12,13,16,24,4,4], 1): pn.column_dimensions[get_column_letter(c)].width = w

# ---------------- LEIA-ME ----------------
ws = wb["Sheet"]; ws.title = "LEIA-ME"
ws.column_dimensions["A"].width = 20; ws.column_dimensions["B"].width = 100
L = [("Acompanhamento de Estudos – Power BI e Excel", None),
     (f"{N} tópicos organizados em 8 semanas, ligados aos 4 pacotes do portfólio e à Base_Ficticia_ECommerce.xlsx.", None),
     ("", None), ("ABA", "COMO USAR"),
     ("Painel", "Visão geral automática: % concluído, horas, progresso por bloco e quais pacotes já podem ser vendidos."),
     ("Checklist", "Lista de tópicos. Altere a coluna Status (lista suspensa) e preencha as datas. Tópicos concluídos ficam riscados."),
     ("Diario", "Registre cada sessão de estudo: data, ID do tópico (lista), horas e o que aprendeu. As horas somam no Checklist e no Painel."),
     ("Cronograma", "Plano de 8 semanas. Altere a data de início (célula C3) e as datas das semanas se ajustam."),
     ("", None), ("CONVENÇÕES", None),
     ("Fundo amarelo", "Células para você preencher (Status, datas, observações, diário)."),
     ("Texto azul", "Valores digitados por você. Texto preto = fórmulas (não altere)."),
     ("Status", "Não iniciado · Estudando · Revisar · Concluído"),
     ("Nível", "Básico (verde) · Intermediário (amarelo) · Avançado (vermelho)"),
     ("Pacotes", "1, 2, 3, 4 = pacotes do portfólio · E = Extras · Todos = vale para todos os pacotes"),
     ("", None), ("DICA", "Marque como Concluído só depois de fazer o Exercício Prático na base de e-commerce. Use 'Revisar' para tópicos que precisam de reforço."),
     ("Novo tópico", "Adicione uma linha no final da tabela do Checklist (ela se expande sozinha) com um novo ID, ex.: T62."),
]
for i, (a, b) in enumerate(L, 1):
    ws[f"A{i}"] = a; ws[f"A{i}"].font = Font(name=F, bold=True)
    if b: ws[f"B{i}"] = b; ws[f"B{i}"].font = Font(name=F); ws[f"B{i}"].alignment = Alignment(wrap_text=True, vertical="top")
ws["A1"].font = Font(name=F, bold=True, size=14, color=NAVY)
ws.merge_cells("A2:B2"); ws["A2"].font = Font(name=F, italic=True)
for c in "AB": ws[f"{c}4"].font = Font(name=F, bold=True, color="FFFFFF"); ws[f"{c}4"].fill = HDR
ws["A10"].font = ws["A16"].font = Font(name=F, bold=True, color=NAVY)
ws["B11"].fill = INPUT; ws["B12"].font = Font(name=F, color="0000FF")

wb.move_sheet("LEIA-ME", offset=-wb.index(wb["LEIA-ME"]) + 1)
wb.active = 0
wb.save(OUT)
print(N, "tópicos")
